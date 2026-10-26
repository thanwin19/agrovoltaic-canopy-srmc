"""Update the A3 poster: final title and the six-sensor controller box.  python update_poster.py  (then export with PowerPoint)"""
import os, shutil, json
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.abspath(os.path.join(HERE, '..', '..')); RES = os.path.join(HERE, 'results')
PD = os.path.join(ROOT, 'Review5_Economics_and_Final', 'documents', 'poster'); src = os.path.join(PD, 'SRMC_Poster_A3.pptx')
bk = os.path.join(PD, '_before_update'); os.makedirs(bk, exist_ok=True)
for f in ['SRMC_Poster_A3.pptx', 'SRMC_Poster_A3.pdf', 'SRMC_Poster_A3.png']:
    if not os.path.exists(os.path.join(bk, f)): shutil.copy(os.path.join(PD, f), bk)
prs = Presentation(os.path.join(bk, 'SRMC_Poster_A3.pptx')); s = prs.slides[0]; sh = list(s.shapes)
n = {r['strategy']: r for r in json.load(open(os.path.join(RES, 'accuracy_summary.json')))['results']}['SRMC-SC (noisy sensors)']
def settext(shape, lines, size=None):
    tf = shape.text_frame; paras = tf.paragraphs
    for i, line in enumerate(lines):
        p = paras[i] if i < len(paras) else tf.add_paragraph()
        if p.runs:
            r0 = p.runs[0]; r0.text = line
            for r in p.runs[1:]: r.text = ''
        else:
            r0 = p.add_run(); r0.text = line
            if paras[0].runs: r0.font.size = paras[0].runs[0].font.size; r0.font.name = paras[0].runs[0].font.name; r0.font.color.rgb = paras[0].runs[0].font.color.rgb
        if size: r0.font.size = Pt(size)
    for p in paras[len(lines):]:
        for r in p.runs: r.text = ''
settext(sh[1], ['Agrovoltaic Canopy for Protection and Sustainable Farming'], 21)
settext(sh[2], ['Six sensors tilt the panels to optimise crop light and electricity; a gutter harvests rain; power is sold to a nearby industry. Digital-twin study, TCE ground, Madurai (9.88 N).'], 12.5)
settext(sh[17], ['THE SIX-SENSOR CONTROLLER'])
settext(sh[18], ['• Six sensors: pyranometer, PAR, DHT11, soil moisture, rain, wind. Rain = flat roof (3 deg); night = park.', '• Otherwise: best tilt 15-60 deg for 50 % electricity + 50 % crop light. PAR corrects the model; heat or dry soil = more shade.', '• Year-long test: %.1f %% of the best score (electricity %.0f %%, crop light %.0f %%).' % (n['J_efficiency_pct'], n['pv_index_pct'], n['crop_index_pct'])], 11)
old = sh[19]; l, t, w, h = old.left, old.top, old.width, old.height; old._element.getparent().remove(old._element)
im = Image.open(os.path.join(RES, 'fig_accuracy_both_goals.png')); sc = min(5.2 / im.size[0], 1.55 / im.size[1]); pw, ph = im.size[0] * sc, im.size[1] * sc
s.shapes.add_picture(os.path.join(RES, 'fig_accuracy_both_goals.png'), Inches(6.0 + (5.3 - pw) / 2), Inches(6.15), Inches(pw), Inches(ph))
settext(sh[20], ['Both goals (top-right is best): six-sensor controller 99.6 % of the best score'], 10)
sh7 = sh[7].text_frame.paragraphs
for p_ in sh7:
    t_ = ''.join(r.text for r in p_.runs)
    if 'hourly controller' in t_: p_.runs[0].text = '• Six sensors tell the controller how to balance electricity and crop light.'; [setattr(r, 'text', '') for r in p_.runs[1:]]
tbl = [x for x in sh if x.has_table][0].table
for r in tbl.rows:
    c = r.cells[0]
    if c.text == 'SRMC-DT controller':
        c.text_frame.paragraphs[0].runs[0].text = 'Hourly plan (SRMC-DT)'
prs.save(src); print('poster pptx saved', round(pw, 2), round(ph, 2))
