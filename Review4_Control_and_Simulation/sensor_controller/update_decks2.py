"""Second pass on the decks: clarify old controller slides (hourly plan = planner) and the wind-stow rows."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
def rep(prs, pairs):
    n = 0
    for s in prs.slides:
        for sh in s.shapes:
            frames = []
            if sh.has_text_frame: frames.append(sh.text_frame)
            if getattr(sh, 'has_table', False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells: frames.append(c.text_frame)
            for tf in frames:
                for para in tf.paragraphs:
                    tx = ''.join(r.text for r in para.runs); new = tx
                    for a, b in pairs: new = new.replace(a, b)
                    if new != tx and para.runs: para.runs[0].text = new; [setattr(r, 'text', '') for r in para.runs[1:]]; n += 1
    return n
p = os.path.join(ROOT, 'Review4_Control_and_Simulation', 'ppt', 'Review4_Control_and_Simulation.pptx'); prs = Presentation(p)
print(rep(prs, [('Simulink model of the controller', 'Simulink model of the hourly-plan controller (planner and fallback)'), ('Simulink test day: rain and wind gust', 'Simulink test day of the hourly plan: rain and wind gust'),
                ('Why hourly control is worth it - and what it does not prove', 'Why sensor-driven control is worth it - and what it does not prove')])); prs.save(p)
p = os.path.join(ROOT, 'Review3_Digital_Twin', 'ppt', 'Review3_Digital_Twin.pptx'); prs = Presentation(p)
print(rep(prs, [('stow at > 15 m/s', 'wind stow at 15 m/s (anemometer)'), ('Simulated sensors and the decision they feed', 'Simulated sensors and the decision they feed (six sensors + modelled panel temperature)')])); prs.save(p)
