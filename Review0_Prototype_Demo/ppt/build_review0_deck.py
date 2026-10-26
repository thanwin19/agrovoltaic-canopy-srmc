"""Builds Review0_Prototype_Demo.pptx from the real result files (python-pptx). Run:  python build_review0_deck.py"""
import os, csv
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__)); RES = os.path.join(os.path.dirname(HERE), 'results')
DEEP = RGBColor(0x14, 0x34, 0x2B); LEAF = RGBColor(0x4E, 0x9A, 0x2E); AMBER = RGBColor(0xE8, 0xA3, 0x17); SKY = RGBColor(0x2D, 0x7F, 0xC1)
INK = RGBColor(0x1B, 0x2A, 0x22); MUTE = RGBColor(0x5B, 0x6B, 0x62); PAPER = RGBColor(0xF6, 0xF7, 0xF2); WHITE = RGBColor(255, 255, 255); RED = RGBColor(0xB0, 0x30, 0x30); PALE = RGBColor(0xE6, 0xEE, 0xE0)
FONT = 'Calibri'; HFONT = 'Cambria'

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5); BLANK = prs.slide_layouts[6]

def bg(slide, color):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = color

def tb(slide, x, y, w, h, paras, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, font=FONT, anchor=MSO_ANCHOR.TOP, italic=False, space=4):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = s.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    if isinstance(paras, str): paras = [paras]
    for i, t in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(space)
        runs = t if isinstance(t, list) else [(t, {})]
        for text, o in runs:
            r = p.add_run(); r.text = text; r.font.size = Pt(o.get('size', size)); r.font.bold = o.get('bold', bold); r.font.italic = o.get('italic', italic)
            r.font.color.rgb = o.get('color', color); r.font.name = o.get('font', font)
    return s

def bullets(slide, x, y, w, h, items, size=16, color=INK, gap=6):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = s.text_frame; tf.word_wrap = True
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(gap)
        runs = t if isinstance(t, list) else [(t, {})]
        first = True
        for text, o in runs:
            r = p.add_run(); r.text = (u'•  ' if first else '') + text; first = False
            r.font.size = Pt(o.get('size', size)); r.font.bold = o.get('bold', False); r.font.color.rgb = o.get('color', color); r.font.name = FONT
    return s

def title(slide, text, sub=None, dark=False):
    tb(slide, 0.6, 0.35, 12.1, 0.8, text, size=(32 if len(text) <= 46 else 27 if len(text) <= 62 else 23), bold=True, color=WHITE if dark else DEEP, font=HFONT, anchor=MSO_ANCHOR.MIDDLE)
    if sub: tb(slide, 0.6, 1.1, 12.1, 0.45, sub, size=16, color=PALE if dark else MUTE)

def box(slide, x, y, w, h, fill=WHITE, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h)); s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0] = radius
    return s

def pic(slide, path, x, y, w, h):
    iw, ih = Image.open(path).size; r = min(w / iw, h / ih); pw, ph = iw * r, ih * r
    return slide.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))

def table(slide, x, y, w, rows, colw, size=12, head=DEEP, rowh=0.36, zebra=True, bold_first=False):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(rowh * len(rows))); t = shp.table
    for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(rowh)
        for j, val in enumerate(row):
            c = t.cell(i, j); c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.02); c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]; r = p.add_run(); r.text = str(val); r.font.size = Pt(size); r.font.name = FONT
            if i == 0: c.fill.solid(); c.fill.fore_color.rgb = head; r.font.bold = True; r.font.color.rgb = WHITE
            else:
                c.fill.solid(); c.fill.fore_color.rgb = PALE if (zebra and i % 2 == 0) else WHITE; r.font.color.rgb = INK; r.font.bold = bold_first and j == 0
                if str(val) in ('PASS', 'REAL'): r.font.color.rgb = LEAF; r.font.bold = True
                if str(val) in ('SIM', 'FAIL'): r.font.color.rgb = AMBER if str(val) == 'SIM' else RED; r.font.bold = True
    return t

def notes(slide, text): slide.notes_slide.notes_text_frame.text = text

def new(dark=False):
    s = prs.slides.add_slide(BLANK); bg(s, DEEP if dark else PAPER); return s

def chip(slide, x, y, w, text, fill=LEAF, size=13, h=0.34):
    b = box(slide, x, y, w, h, fill=fill, radius=0.5); tf = b.text_frame; tf.margin_left = tf.margin_right = Inches(0.05); p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return b

def read_csv(path):
    with open(path, newline='', encoding='utf-8-sig') as f: return list(csv.DictReader(f))

tests = read_csv(os.path.join(RES, 'test_results.csv')); day = read_csv(os.path.join(RES, 'day_comparison.csv'))
npass = sum(1 for t in tests if t['result'] == 'PASS')

# ============ 1 title ============
s = new(dark=True)
pic(s, os.path.join(RES, 'proof', 'hero.png'), 6.9, 1.3, 6.1, 4.9)
tb(s, 0.6, 1.2, 6.2, 0.5, 'REVIEW 0', size=18, bold=True, color=AMBER)
tb(s, 0.6, 1.7, 6.2, 2.6, 'Sensor-driven tilt control that co-optimises solar energy and crop light', size=34, bold=True, color=WHITE, font=HFONT)
tb(s, 0.6, 4.55, 6.1, 1.2, ['SRMC: wire-driven agrivoltaic canopy', 'Prototype sensors + MATLAB digital twin  |  TCE ground, Madurai'], size=17, color=PALE)
tb(s, 0.6, 6.0, 6.2, 1.0, ['Student: ______________________', 'Guide: ______________________     Date: __________'], size=15, color=PALE)
notes(s, 'Review 0 shows the live chain: real sensors on the ESP board -> MATLAB -> tilt decision. Say at the start: pyranometer, PAR and wind are SIMULATED sensors; DHT11, soil and rain are REAL.')

# ============ 2 objective ============
s = new(); title(s, 'What the system decides', 'One decision every few seconds: the panel angle that gives energy AND the light the crop needs')
cards = [('Pyranometer', 'how much sunlight, so how much energy to harvest', SKY, 'SIM'), ('PAR sensor', 'light actually reaching each crop zone', LEAF, 'SIM'), ('Soil moisture', 'root-zone water: irrigation and stress', RGBColor(0x8B, 0x5A, 0x2B), 'REAL'),
         ('DHT11', 'temperature and humidity: VPD (plant water stress)', AMBER, 'REAL'), ('Rain sensor', 'rain: panels to a 3 degree slope, water to the cistern', RGBColor(0x2D, 0x5F, 0x9A), 'REAL'), ('Wind', 'gust above 15 m/s: panels stowed edge-on (90 degrees)', MUTE, 'SIM')]
for i, (h, d, col, tag) in enumerate(cards):
    x = 0.6 + (i % 3) * 4.15; y = 1.85 + (i // 3) * 2.15
    box(s, x, y, 3.9, 1.95, fill=WHITE, line=PALE); box(s, x + 0.2, y + 0.2, 0.55, 0.55, fill=col, shape=MSO_SHAPE.OVAL)
    tb(s, x + 0.9, y + 0.17, 2.3, 0.5, h, size=20, bold=True, color=DEEP); chip(s, x + 3.0, y + 0.25, 0.75, tag, fill=LEAF if tag == 'REAL' else AMBER, size=11, h=0.3)
    tb(s, x + 0.25, y + 0.9, 3.45, 0.95, d, size=15, color=INK)
tb(s, 0.6, 6.25, 12.1, 0.8, [[('Goal: ', {'bold': True, 'color': DEEP}), ('a tilt rule that is better for energy AND crops than pure sun tracking or a fixed angle, shown with numbers.', {})]], size=17)
notes(s, 'Six inputs, one output (the angle). The controller is a rule stack plus an optimiser. Real vs simulated is labelled on every card.')

# ============ 3 real vs sim ============
s = new(); title(s, 'What is real and what is simulated', 'Stated openly, so every number on the next slides can be trusted for what it is')
rows = [['Quantity', 'Source on the day', 'Status', 'Why'],
        ['Air temperature, humidity', 'DHT11 on the prototype', 'REAL', 'sensor installed'],
        ['Soil moisture', 'soil moisture module -> A0', 'REAL', 'sensor installed, calibrated dry/wet'],
        ['Rain', 'rain sensor module -> D5', 'REAL', 'sensor installed'],
        ['Pyranometer (W/m2)', 'MATLAB sun model for TCE + sky icon (clear, hazy, partly cloudy, overcast)', 'SIM', 'no pyranometer available yet'],
        ['PAR above / under canopy', 'computed from the pyranometer value and the canopy shadow model', 'SIM', 'no PAR sensor available yet'],
        ['Wind speed', 'slider on the dashboard', 'SIM', 'no anemometer yet'],
        ['Panel tilt', 'shown on the MATLAB drawing only', 'SIM', 'no actuator on the prototype']]
table(s, 0.6, 1.75, 12.1, rows, [2.7, 5.3, 1.1, 3.0], size=14, rowh=0.52)
box(s, 0.6, 6.05, 12.1, 0.95, fill=WHITE, line=AMBER)
tb(s, 0.8, 6.1, 11.7, 0.85, 'The simulated sensors use the same equations as the validated model (PVGIS comparison +5.0 %, Blender ray tracing 0.010). They are placeholders for a real pyranometer and PAR sensor, which would only replace two inputs.', size=15, color=INK, anchor=MSO_ANCHOR.MIDDLE)
notes(s, 'Judges will ask about the pyranometer. Answer: it is a software sensor for Review 0, labelled SIM on the screen. The control code does not change when a real one is connected.')

# ============ 4 architecture ============
s = new(); title(s, 'System architecture', 'Sensors -> ESP board -> USB -> MATLAB -> decision -> tilt on screen and irrigation')
def node(x, y, w, h, head, sub, fill, fg=WHITE):
    b = box(s, x, y, w, h, fill=fill); tb(s, x + 0.05, y + 0.1, w - 0.1, 0.5, head, size=17, bold=True, color=fg, align=PP_ALIGN.CENTER); tb(s, x + 0.08, y + 0.6, w - 0.16, h - 0.65, sub, size=13, color=fg, align=PP_ALIGN.CENTER)
def arrow(x1, y1, x2, y2):
    c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); c.line.color.rgb = MUTE; c.line.width = Pt(2.5)
    ln = c.line._get_or_add_ln(); from lxml import etree; te = etree.SubElement(ln, '{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd'); te.set('type', 'triangle')
node(0.5, 2.0, 2.6, 2.6, 'REAL sensors', 'DHT11 (temperature, humidity)\nsoil moisture\nrain sensor', LEAF)
node(3.7, 2.0, 2.3, 2.6, 'ESP8266 / ESP32', 'reads sensors every second\nsends one text line over USB (115200 baud)', SKY)
node(6.6, 2.0, 2.8, 2.6, 'MATLAB dashboard', 'parses the line, calibrates soil\nlogs CSV, draws live plots', DEEP)
node(10.0, 2.0, 2.8, 2.6, 'Controller', 'wind stow > rain 3 deg > night park > optimise energy + crop light', AMBER, INK)
node(3.7, 5.2, 2.3, 1.5, 'SIM sensors', 'pyranometer, PAR, wind (sliders + sky icon)', RGBColor(0x8A, 0x6A, 0x20))
node(10.0, 5.2, 2.8, 1.5, 'Outputs', 'panel angle on screen, pump ON/OFF, cistern level, evidence files', MUTE)
arrow(3.1, 3.3, 3.7, 3.3); arrow(6.0, 3.3, 6.6, 3.3); arrow(9.4, 3.3, 10.0, 3.3); arrow(4.85, 5.2, 7.4, 4.6); arrow(11.4, 4.6, 11.4, 5.2)
tb(s, 0.5, 5.25, 3.0, 1.4, ['Backup: ESP32 with the same serial protocol', 'MATLAB also runs without hardware (Simulate mode)'], size=14, color=MUTE)
notes(s, 'USB serial is used because it needs no WiFi and is the most reliable in a lab. The same MATLAB code runs in Simulate mode, so the demo works even if a wire comes loose.')

# ============ 5 wiring ============
s = new(); title(s, 'Hardware wiring', 'NodeMCU ESP8266 (CP2102) main board, ESP32 as backup; everything on 3.3 V and one common ground')
rows = [['Sensor', 'Sensor pin', 'ESP8266 NodeMCU', 'ESP32'], ['DHT11', 'VCC / DATA / GND', '3V3 / D2 / GND', '3V3 / GPIO4 / GND'], ['Soil moisture', 'VCC / AO / GND', '3V3 / A0 / GND', '3V3 / GPIO34 / GND'], ['Rain sensor', 'VCC / DO / GND', '3V3 / D5 / GND', '3V3 / GPIO27 / GND (AO to GPIO35 optional)']]
table(s, 0.6, 1.75, 12.1, rows, [2.3, 2.9, 3.1, 3.8], size=15, rowh=0.55, bold_first=True)
bullets(s, 0.6, 4.2, 6.1, 2.9, ['ESP8266 has only one analog pin: soil on A0, rain on the digital output D5', 'D4 on-board LED lights while rain is detected', 'Avoid D3, D4 and D8 for sensors (they set the boot mode)', 'Never put 5 V on A0 or a D pin'], size=16)
box(s, 7.0, 4.2, 5.7, 2.35, fill=WHITE, line=PALE); tb(s, 7.2, 4.3, 5.3, 0.4, 'Bring-up check (5 minutes)', size=17, bold=True, color=DEEP)
bullets(s, 7.2, 4.75, 5.4, 2.2, ['Upload the sketch, open the Serial Monitor at 115200', 'See SRMC,5231,31.0,54.0,712,0,-1 every second', 'Wet the rain plate: rainWet turns 1', 'Push the soil probe into wet soil: soilRaw drops'], size=14, gap=4)
notes(s, 'Full pin map is in firmware/PIN_MAP.md. The photo of the board shows the CP2102 USB chip, so install that driver if no COM port appears.')

# ============ 6 firmware ============
s = new(); title(s, 'Firmware and data protocol', 'One text line per second; the same MATLAB code reads either board')
box(s, 0.6, 1.8, 12.1, 1.3, fill=DEEP); tb(s, 0.85, 1.95, 11.6, 1.0, ['SRMC,<millis>,<tempC>,<rh>,<soilRaw>,<rainWet>,<rainRaw>', 'SRMC,5231,31.0,54.0,712,0,-1'], size=22, color=WHITE, font='Consolas')
bullets(s, 0.6, 3.4, 6.0, 3.6, ['DHT11 decoded in the sketch itself (no library to install); a failed read sends -999 and MATLAB keeps the last good value', 'Soil: 8 reads averaged, then calibrated to % with two captured points (dry in air, wet in water)', 'Rain: digital output, sensitivity set by the module potentiometer'], size=16)
bullets(s, 6.9, 3.4, 5.8, 3.6, ['Parser rejects boot messages and noise (tested: T01 to T03)', 'Source = Serial uses the real board; Source = Simulate generates realistic data for rehearsal', 'Every run can be saved: dashboard image + CSV log'], size=16)
notes(s, 'The firmware was written without access to the board, so the bench checklist in the documents folder is the first thing to run on the day.')

# ============ 7 controller ============
s = new(); title(s, 'Control algorithm (simple enough to explain, precise enough to test)', 'Rules in priority order, then an optimiser over every angle from 15 to 60 degrees in 0.5 degree steps')
steps = [('1  Wind >= 15 m/s', 'stow at 90 degrees', RED), ('2  Rain sensor wet', 'slope 3 degrees to the gutter, fill the cistern', SKY), ('3  Night', 'park at 15 degrees', MUTE), ('4  Otherwise', 'optimise energy + crop light', LEAF)]
for i, (h, d, col) in enumerate(steps):
    x = 0.6 + i * 3.08; box(s, x, 1.85, 2.9, 1.55, fill=col); tb(s, x + 0.1, 1.95, 2.7, 0.5, h, size=18, bold=True, color=WHITE); tb(s, x + 0.1, 2.5, 2.7, 0.85, d, size=15, color=WHITE)
box(s, 0.6, 3.75, 7.0, 3.2, fill=WHITE, line=PALE)
tb(s, 0.8, 3.85, 6.6, 0.45, 'Optimiser (rule 4)', size=18, bold=True, color=DEEP)
tb(s, 0.8, 4.35, 6.6, 0.55, 'J(a) = w * PV(a) / PVmax + (1 - w) * Crop(a) / Cropmax', size=19, bold=True, color=INK, font='Consolas')
bullets(s, 0.8, 5.0, 6.6, 1.95, ['PV(a): module power for angle a (row shading, incidence angle, cell temperature from the DHT11)', 'Crop(a): mean light factor of the 4 crop zones, from PAR under / above the canopy', 'w = 0.5 by default; the slider moves it from crop-only to energy-only'], size=14, gap=3)
box(s, 7.9, 3.75, 4.8, 3.2, fill=WHITE, line=PALE); tb(s, 8.1, 3.85, 4.4, 0.45, 'Stress makes the controller smarter', size=18, bold=True, color=DEEP)
bullets(s, 8.1, 4.4, 4.45, 2.5, ['VPD from DHT11 > 2.5 kPa or dry soil raises the stress index', 'Stress lowers the light the crops need (by up to 30 %), so the controller can accept more shade and more energy', 'Pump: ON below 35 % soil moisture, OFF above 55 % (hysteresis), water from the rain cistern'], size=14, gap=3)
notes(s, 'Explain in one breath: safety first (wind), then water (rain), then the optimiser. The optimiser is a plain search over 91 angles with a transparent score, not a black box.')

# ============ 8 dashboard ============
s = new(); title(s, 'Live MATLAB dashboard', 'Same code and screen for Simulate and Serial mode; two situations shown large')
for i, (f, cap) in enumerate([('dashboard_heat_stress.png', 'Hot, dry air (VPD 4.6 kPa): stress index 1.0, controller chooses 43 degrees at 13 h'), ('dashboard_rain_harvest.png', 'Rain sensor wet: panels to the 3 degree slope, cistern fills')]):
    x = 0.6 + i * 6.15; box(s, x, 1.7, 5.95, 3.55, fill=WHITE, line=PALE); pic(s, os.path.join(RES, 'screens', f), x + 0.05, 1.75, 5.85, 3.2); tb(s, x + 0.1, 4.93, 5.75, 0.3, cap, size=12, color=MUTE, align=PP_ALIGN.CENTER)
bullets(s, 0.6, 5.5, 12.1, 1.6, ['Left: controls (source, sky icon, wind, calibration, weight).  Centre: canopy drawing with sun and sky icon, panel angle, decision text with REAL / SIM tags.  Right: live plots.', 'Other situations are in results/screens: clear noon (43.5 degrees), gust 18 m/s (stow at 90 degrees), dry soil (pump ON), low morning sun (60 degrees).'], size=14, gap=5)
notes(s, 'Left column: controls. Centre: canopy drawing with sun icon, sky icon and panel angle. Right: live plots. Real and simulated values are tagged in the text panel.')

# ============ 9 tests ============
s = new(); title(s, 'Test evidence 1: decision logic (%d of %d software-in-the-loop tests pass)' % (npass, len(tests)), 'Simulated inputs; this proves the rules and the physics, not the sensors (sensor tests are done on the bench)')
short = {'T01': 'Parse a good serial line', 'T02': 'DHT11 failure (-999) flagged, not used', 'T03': 'Boot / garbage line ignored', 'T04': 'Soil calibration 0 / 50 / 100 %', 'T05': 'VPD at 30 C, 50 % RH = 2.12 kPa', 'T06': 'Rain wet -> 3 degree slope', 'T07': 'Wind 15.0 stows, 14.9 does not', 'T08': 'Wind stow overrides rain',
         'T09': 'Night: park 15 degrees, no PV', 'T10': 'Optimiser = brute-force best angle', 'T11': 'Noon PV >= fixed 45 degrees', 'T12': 'Heat stress changes the angle', 'T13': 'Overcast GHI < 40 % of clear', 'T14': 'Pump hysteresis 35 / 55 %', 'T15': 'Empty cistern stops the pump', 'T16': 'Equals the validated engine'}
half = (len(tests) + 1) // 2
for k, part in enumerate([tests[:half], tests[half:]]):
    rows = [['ID', 'Test', 'Result']] + [[t['id'], short.get(t['id'], t['test']), t['result']] for t in part]
    table(s, 0.6 + k * 6.15, 1.75, 5.95, rows, [0.7, 4.4, 0.85], size=13, rowh=0.5)
tb(s, 0.6, 6.55, 12.1, 0.6, 'Files: results/test_results.csv and .md; reproduce with review0_run_tests in MATLAB (about 1 minute).', size=14, color=MUTE)
notes(s, 'T16 is the strongest proof: the MATLAB functions give the same zone light and PV as the JavaScript engine used in the twin. T10 checks the optimiser against brute force.')

# ============ 10 result 1 chart ============
s = new(); title(s, 'Result: energy and crop light together (clear day, simulated sky)', 'Each strategy as a percentage of the best strategy for that quantity; day 105, TCE, 6 to 18 h')
clear = [r for r in day if r['condition'] == 'Clear day']
cd = CategoryChartData(); cd.categories = [r['strategy'] for r in clear]
cd.add_series('PV energy (% of best)', [round(float(r['pv_pct_of_best']), 1) for r in clear]); cd.add_series('Crop light (% of best)', [round(float(r['crop_pct_of_best']), 1) for r in clear])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.7), Inches(8.2), Inches(5.3), cd); ch = gf.chart
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(13)
ch.plots[0].gap_width = 60; ch.plots[0].has_data_labels = True; dl = ch.plots[0].data_labels; dl.font.size = Pt(11); dl.number_format = '0'; dl.number_format_is_linked = False; dl.position = XL_LABEL_POSITION.OUTSIDE_END
ch.series[0].format.fill.solid(); ch.series[0].format.fill.fore_color.rgb = AMBER; ch.series[1].format.fill.solid(); ch.series[1].format.fill.fore_color.rgb = LEAF
ch.category_axis.tick_labels.font.size = Pt(11); ch.value_axis.tick_labels.font.size = Pt(11); ch.value_axis.maximum_scale = 115; ch.value_axis.has_major_gridlines = True; ch.value_axis.major_gridlines.format.line.color.rgb = PALE
def getv(name, key): return float([r for r in clear if r['strategy'] == name][0][key])
sr = 'SRMC controller (50/50)'; tr = 'East-west sun tracking'; f45 = 'Fixed 45 deg'
box(s, 9.1, 1.8, 3.6, 5.1, fill=WHITE, line=PALE); tb(s, 9.25, 1.9, 3.3, 0.45, 'Reading the chart', size=18, bold=True, color=DEEP)
bullets(s, 9.25, 2.4, 3.35, 4.4, [[('Sun tracking: ', {'bold': True}), ('best energy (100 %%) but only %.0f %% of the crop light' % getv(tr, 'crop_pct_of_best'), {})],
                                   [('Fixed 45 degrees: ', {'bold': True}), ('%.0f %% energy, %.0f %% crop' % (getv(f45, 'pv_pct_of_best'), getv(f45, 'crop_pct_of_best')), {})],
                                   [('SRMC controller: ', {'bold': True}), ('%.0f %% energy, %.0f %% crop; average %.0f %% (tracking %.0f %%, fixed 45: %.0f %%)' % (getv(sr, 'pv_pct_of_best'), getv(sr, 'crop_pct_of_best'), getv(sr, 'combined_pct'), getv(tr, 'combined_pct'), getv(f45, 'combined_pct')), {})]], size=14, gap=7)
notes(s, 'Honest reading: the controller does not beat tracking on energy; it gives up about a quarter of the energy to give the crops most of their light. The average favours the controller by construction (equal weights), so the weight slider matters. Pure tracking is outside the wire-drive specification (two positions), shown only as a reference. Numbers come from day_comparison.csv.')

# ============ 11 result 2 ============
s = new(); title(s, 'Result: the angle the controller chooses through the day', 'Compared with pure sun tracking (dashed), clear, partly cloudy and hot-dry days')
pic(s, os.path.join(RES, 'fig_angle_through_day.png'), 0.6, 1.7, 7.9, 5.3)
bullets(s, 8.8, 1.9, 3.9, 5.0, ['Morning: steep (about 55 to 60 degrees) so crops get light while the sun is low', 'Afternoon on clear days: flatter, more energy', 'Partly cloudy day: stays steeper because the crop needs the light that is left', 'Hot, dry air (VPD stress): flatter around midday: more shade for the crop, more energy', 'Tracking would turn negative (west); the wire-drive design has two positions, so it is only a reference'], size=15, gap=8)
notes(s, 'The jumps between half-hour steps are caused by a flat objective; in the real motor a 2 degree dead-band would stop needless movement (listed as next step).')

# ============ 12 result 3 ============
s = new(); title(s, 'Result: one slider moves the farm between "crop first" and "energy first"', 'Noon, clear sky: the energy weight w from 0 to 100 %')
pic(s, os.path.join(RES, 'fig_weight_sweep_noon.png'), 0.6, 1.7, 7.9, 5.2)
wsr = read_csv(os.path.join(RES, 'weight_sweep_noon.csv')); a0 = wsr[0]; a1 = wsr[-1]
rows = [['w (energy)', 'angle', 'PV W/m2', 'crop score']] + [[w['energy_weight_pct'] + ' %', w['angle_deg'] + ' deg', '%.0f' % float(w['pv_Wm2']), '%.2f' % float(w['crop_score'])] for w in wsr[::2]]
table(s, 8.8, 1.9, 3.9, rows, [1.0, 0.95, 1.0, 0.95], size=13, rowh=0.5)
tb(s, 8.8, 5.6, 3.9, 1.3, 'The farmer or operator chooses w; the code does the rest. Review 4 prices the choice.', size=14, color=MUTE)
notes(s, 'This is the practical meaning of co-optimisation: the weight is a policy knob, the optimiser finds the best angle for that policy at every moment.')

# ============ 13 validation ============
s = new(); title(s, 'The model under the demo is validated', 'Same equations as the earlier reviews; independent references, not field data')
for i, (f, cap) in enumerate([('val_pv_vs_pvgis.png', 'PV energy vs PVGIS: +5.0 % (of which +5.2 % is the sunshine dataset)'), ('val_shading_scatter.png', 'Shadow model vs Blender ray tracing: mean difference 0.010'), ('val_structural_utilisation.png', 'Structure: 12.5 m bays fail, the 5 m rope grid passes (screening)')]):
    x = 0.6 + i * 4.1; box(s, x, 1.75, 3.9, 4.1, fill=WHITE, line=PALE); pic(s, os.path.join(RES, 'proof', f), x + 0.1, 1.85, 3.7, 3.0); tb(s, x + 0.15, 4.95, 3.6, 0.85, cap, size=14, color=INK)
tb(s, 0.6, 6.1, 12.1, 0.9, ['JavaScript twin vs MATLAB engine: identical results (0.0000 % on 22 quantities). This checks the code, not the physics.', 'Not done: comparison with measured field data (no pyranometer or field plot yet).'], size=15, color=MUTE)
notes(s, 'These are in the earlier review folders (Review 2 and Review 1). Say clearly what they are not: they are not field measurements.')

# ============ 14 bench evidence ============
s = new(); title(s, 'Bench evidence from the prototype (added on the day)', 'Photos and readings taken at the college; the table is filled with measured values')
for i, cap in enumerate(['Photo: ESP8266 wired to DHT11, soil and rain sensors', 'Screenshot: Serial Monitor showing SRMC lines', 'Screenshot: MATLAB dashboard with the real board (Source = Serial)']):
    x = 0.6 + i * 4.1; b = box(s, x, 1.75, 3.9, 2.6, fill=WHITE, line=MUTE); b.line.dash_style = 4; tb(s, x + 0.2, 2.5, 3.5, 1.2, cap, size=15, color=MUTE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rows = [['Bench test', 'Reference', 'Measured', 'Pass criterion'], ['DHT11 temperature', 'room thermometer / phone', '_____ C vs _____ C', 'within 2 C (datasheet)'], ['DHT11 humidity', 'reference hygrometer', '_____ % vs _____ %', 'within 5 % RH (datasheet)'],
        ['Soil: dry (air) / wet (water)', 'raw readings', 'dry _____  wet _____', 'wet raw clearly below dry raw'], ['Rain detection', 'drop on the plate', 'detected in _____ s', 'rain = 1 within 2 s'], ['Rain dry again', 'wipe the plate', 'cleared in _____ s', 'rain = 0 within 30 s'],
        ['End to end', 'wet plate -> MATLAB', '3 degree mode in _____ s', 'under 5 s']]
table(s, 0.6, 4.55, 12.1, rows, [3.0, 3.0, 3.3, 2.8], size=13, rowh=0.36)
notes(s, 'Replace the three frames with photos and screenshots on the morning of the review; fill the table with the real numbers. Do not leave blanks on screen.')

# ============ 15 feasibility ============
s = new(); title(s, 'Feasibility and practical constraints', 'What limits the prototype today and what the plan is')
rows = [['Constraint', 'Effect', 'Plan'], ['DHT11 accuracy: +-2 C, +-5 % RH', 'VPD estimate is approximate', 'use as a trend and a stress flag; upgrade to SHT31 later'],
        ['Resistive soil probe corrodes and drifts', 'moisture % drifts over weeks', 'capacitive probe; recalibrate (two-point) each session'], ['ESP8266 has one analog input', 'soil uses A0, rain uses the digital pin', 'ESP32 backup (several ADC pins)'],
        ['No real pyranometer / PAR sensor', 'light inputs are simulated', 'add a calibrated silicon pyranometer and a quantum PAR sensor; no code change'], ['No actuator on the prototype', 'tilt shown on screen only', 'wire-drive winch with 0.5 degree encoder (full-scale design is in Reviews 1 and 4)'],
        ['USB cable to a PC', 'not field deployable', 'WiFi/MQTT version of the same protocol'], ['Screening-level models', 'results are estimates', 'validate against local weather data and a field plot']]
table(s, 0.6, 1.75, 12.1, rows, [3.7, 3.6, 4.8], size=14, rowh=0.62)
notes(s, 'Feasibility is about being honest on constraints and showing a path. Every constraint has a cheap mitigation.')

# ============ 16 SDG and cost ============
s = new(); title(s, 'SDG mapping and cost-effectiveness', 'Impact of the idea (left) and the cost of the prototype kit (right)')
rows = [['SDG', 'How the project contributes', 'Indicator'], ['7 Clean energy', 'solar power from farmland, sold to industry', 'kWh per hectare'], ['2 Zero hunger', 'crops keep the light they need under the panels', 'crop-light score'],
        ['6 Water', 'rain harvested from the panel roof, drip irrigation', 'm3 harvested; % rain-fed'], ['13 Climate action', 'avoided grid emissions', 'kg CO2 per kWh'], ['1 / 8 Livelihood', 'monthly income from the industry plus crops', 'INR per month']]
table(s, 0.6, 1.75, 6.9, rows, [1.7, 3.5, 1.7], size=13, rowh=0.62)
rows = [['Prototype part', 'Indicative price', 'Your bill'], ['NodeMCU ESP8266', 'INR 250-400', 'INR ____'], ['ESP32 (backup)', 'INR 400-700', 'INR ____'], ['DHT11', 'INR 60-150', 'INR ____'], ['Soil moisture module', 'INR 50-120', 'INR ____'], ['Rain sensor module', 'INR 50-120', 'INR ____'], ['Wires, breadboard, USB cable', 'INR 100-250', 'INR ____']]
table(s, 7.8, 1.75, 4.9, rows, [2.4, 1.4, 1.1], size=13, rowh=0.5)
box(s, 0.6, 5.55, 12.1, 1.4, fill=WHITE, line=AMBER)
tb(s, 0.8, 5.6, 11.7, 1.3, ['Indicative prices are market ranges for hobby-grade parts; replace them with your bill. The complete sensor kit costs a few hundred to about one thousand rupees.', 'A calibrated pyranometer costs far more (typically tens of thousands of rupees), which is why Review 0 uses a virtual one and a low-cost silicon sensor is the planned upgrade.'], size=14, color=INK, anchor=MSO_ANCHOR.MIDDLE)
notes(s, 'Prices are indicative and must be replaced by the real bill before presenting a cost figure. The economic case of the full system (farmer income, payback) is in Review 5 and depends on assumed prices.')

# ============ 17 progress ============
s = new(); title(s, 'Time management and progress', 'The computer study is complete; Review 0 adds the real sensors')
phases = [('Design + CAD', 'done', LEAF), ('Solar + shadow', 'done', LEAF), ('Digital twin', 'done', LEAF), ('Controller', 'done', LEAF), ('Economics', 'done', LEAF), ('Review 0 sensors', 'now', AMBER), ('Real PAR sensor', 'next', MUTE), ('Actuator + field', 'next', MUTE)]
for i, (h, st, col) in enumerate(phases):
    x = 0.45 + i * 1.57; c = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(2.0), Inches(1.7), Inches(1.3)); c.fill.solid(); c.fill.fore_color.rgb = col; c.line.fill.background(); c.adjustments[0] = 0.2
    tf = c.text_frame; tf.word_wrap = True; tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.1); p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = h; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = WHITE if col != AMBER else INK; r.font.name = FONT
    tb(s, x + 0.1, 3.4, 1.5, 0.4, {'done': 'completed', 'now': 'this review', 'next': 'planned'}[st], size=13, color=MUTE, align=PP_ALIGN.CENTER)
rows = [['Milestone', 'Evidence file'], ['Design preset, structural screening, CAD drawings, BOM', 'Review 1 folder'], ['Sun, PV and shadow model validated (PVGIS, Blender)', 'Review 2 folder'], ['Digital twin, JavaScript vs MATLAB check', 'Review 3 folder'],
        ['Controller, Simulink test day, comparisons', 'Review 4 folder'], ['Sensitivity, Monte Carlo, final report, poster', 'Review 5 folder'], ['Review 0: firmware, dashboard, 16 tests, results', 'Review0 folder']]
table(s, 0.6, 4.1, 12.1, rows, [8.0, 4.1], size=14, rowh=0.4)
notes(s, 'Each earlier review has its own folder with deck, results and README; use them to answer questions about any stage.')

# ============ 18 summary ============
s = new(dark=True); title(s, 'Summary and next steps', dark=True)
bullets(s, 0.6, 1.6, 7.2, 5.2, ['Real DHT11, soil and rain sensors feed MATLAB through an ESP board; pyranometer, PAR and wind are simulated and labelled', 'One controller: wind stow, rain slope, night park, then an optimiser for energy + crop light with stress adaptation', '%d of %d logic tests pass; the physics equals the validated engine' % (npass, len(tests)), 'Clear day: the controller keeps %.0f %% of the best energy and %.0f %% of the best crop light; sun tracking keeps 100 %% and %.0f %%' % (getv(sr, 'pv_pct_of_best'), getv(sr, 'crop_pct_of_best'), getv(tr, 'crop_pct_of_best'))], size=18, color=WHITE, gap=12)
box(s, 8.2, 1.6, 4.5, 3.5, fill=RGBColor(0x1F, 0x4A, 0x3C)); tb(s, 8.4, 1.7, 4.1, 0.5, 'Next', size=22, bold=True, color=AMBER)
bullets(s, 8.4, 2.3, 4.1, 4.1, ['Add a real pyranometer and PAR sensor', 'Add a 2 degree dead-band and a motor driver', 'Log a week of real sensor data', 'Compare against local weather data'], size=17, color=WHITE, gap=10)
tb(s, 0.6, 6.7, 12.1, 0.5, 'Simulation and prototype sensors only. Prices and farmer-income figures are assumptions (see Review 5).', size=14, color=PALE)
notes(s, 'Close with what is claimed and what is not. Invite questions on any slide; the proof index in Review 5 maps every claim to a file.')

out = os.path.join(HERE, 'Review0_Prototype_Demo.pptx'); prs.save(out); print('saved', out, len(prs.slides), 'slides')
