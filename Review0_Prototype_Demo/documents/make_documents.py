"""Creates the supervisor one-page write-up (docx), bench checklist, demo script + Q&A and the log template. Run: python make_documents.py"""
import os, csv
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)

# ------------------------------------------------------------------ one-page write-up
d = Document()
sec = d.sections[0]; sec.top_margin = sec.bottom_margin = Cm(1.5); sec.left_margin = sec.right_margin = Cm(1.8)
st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
def para(text, bold=False, size=None, color=None, after=3, align=None, italic=False):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor(*color)
    if align: p.alignment = align
    return p
def bullet(text):
    p = d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(1); p.add_run(text); return p
para('Review 0: one-page write-up for supervisor approval', bold=True, size=15, color=(0x14, 0x34, 0x2B), after=2)
para('Project: Wire-driven agrivoltaic canopy with sensor-driven tilt control (SRMC), TCE ground, Madurai', bold=True, after=1)
para('Student: ____________________     Register no.: ______________     Guide: ____________________     Date: ____________', after=6)
para('1. Objective', bold=True, color=(0x4E, 0x9A, 0x2E), after=1)
para('Choose the panel tilt angle that gives the most solar energy while the crop below still receives the light it needs. One controller uses: pyranometer (energy), PAR (crop light), soil moisture (irrigation and stress), DHT11 (temperature, humidity, VPD stress), rain sensor (rain-harvest slope) and wind (safety stow).', after=5)
para('2. What will be demonstrated in Review 0', bold=True, color=(0x4E, 0x9A, 0x2E), after=1)
bullet('Real sensors on the prototype: DHT11, soil moisture and rain sensor connected to an ESP8266 NodeMCU (ESP32 as backup); data sent over USB serial to MATLAB every second.')
bullet('Virtual (simulated) pyranometer, PAR and wind in MATLAB, because these sensors are not yet available. They are computed from the validated sun and canopy-shadow model and are labelled SIM on screen.')
bullet('Live decision: rain detected -> panels to a 3 degree slope (harvest); wind >= 15 m/s -> stow at 90 degrees; night -> park; otherwise an optimiser picks the angle (15 to 60 degrees, 0.5 degree steps) that balances energy and crop light. The tilt is shown on the MATLAB screen only (no actuator).')
para('3. Evidence that will be shown', bold=True, color=(0x4E, 0x9A, 0x2E), after=1)
bullet('16 software-in-the-loop tests of the decision logic, all passing; physics identical to the earlier validated model (PVGIS comparison +5.0 %, Blender ray-tracing check 0.010).')
bullet('Bench measurements on the day: DHT11 against a reference thermometer, soil dry/wet calibration, rain detection time, end-to-end response, with photos and serial/MATLAB screenshots and a saved CSV log.')
bullet('Day simulation: the controller keeps about 75 % of the best energy and 89 % of the best crop light on a clear day; pure sun tracking keeps 100 % energy but 27 % of the crop light.')
para('4. Limits stated openly', bold=True, color=(0x4E, 0x9A, 0x2E), after=1)
para('Pyranometer, PAR and wind are simulated; tilt is not mechanised; DHT11 accuracy is +-2 C and +-5 % RH; models are screening-level and validated against PVGIS and Blender, not against field data. Prices and income figures elsewhere in the project are assumptions.', after=5)
para('5. Resources', bold=True, color=(0x4E, 0x9A, 0x2E), after=1)
para('NodeMCU ESP8266 (ESP32 spare), DHT11, soil moisture module, rain sensor module, laptop with MATLAB and Arduino IDE.', after=10)
t = d.add_table(rows=2, cols=3); t.alignment = WD_TABLE_ALIGNMENT.CENTER
for j, h in enumerate(['Student signature', 'Supervisor signature and date', 'HDEE']):
    c = t.cell(1, j); c.text = ''; t.cell(0, j).text = h; t.cell(0, j).paragraphs[0].runs[0].bold = True
for r in t.rows: r.height = Cm(1.6)
d.save(os.path.join(HERE, 'Review0_Supervisor_OnePager.docx'))

# ------------------------------------------------------------------ bench checklist
chk = """# Review 0 - morning checklist (do in this order)

Bring: laptop (MATLAB + Arduino IDE), USB cable, NodeMCU ESP8266 (+ ESP32 spare), DHT11, soil module, rain module, 6+ jumper wires, a cup of water, a dry cloth, a reference thermometer (phone app is acceptable), phone for photos, the printed one-pager for the signature.

## A. Wire (15 min) - details in firmware/PIN_MAP.md
- [ ] DHT11: VCC -> 3V3, DATA -> D2, GND -> GND
- [ ] Soil: VCC -> 3V3, AO -> A0, GND -> GND
- [ ] Rain: VCC -> 3V3, DO -> D5, GND -> GND
- [ ] Photo of the wiring (needed on slide 14)

## B. Upload and check (10 min)
- [ ] Arduino IDE: board "NodeMCU 1.0 (ESP-12E Module)", correct COM port (install CP2102 driver if no port)
- [ ] Open firmware/SRMC_ESP8266_NodeMCU/SRMC_ESP8266_NodeMCU.ino -> Upload
- [ ] Serial Monitor at 115200: lines like SRMC,5231,31.0,54.0,712,0,-1 appear once per second
- [ ] Screenshot of the Serial Monitor (slide 14)
- [ ] Close the Serial Monitor (MATLAB needs the port)
- [ ] If the ESP8266 misbehaves, repeat with the ESP32 sketch and pins (PIN_MAP.md, part B)

## C. Calibrate and measure (25 min) - write the numbers in test_data/bench_log_template.csv
- [ ] DHT11 vs thermometer: ______ C vs ______ C (must be within 2 C); humidity if you have a reference
- [ ] Soil DRY (probe in air): raw = ______      Soil WET (probe in water): raw = ______   (wet must be clearly below dry)
- [ ] In the dashboard: press Capture DRY with the probe in air, then Capture WET with the probe in water
- [ ] Rain: drop water on the plate -> time to rainWet = 1: ______ s; wipe it -> time to 0: ______ s
- [ ] If rain never turns 1: turn the blue potentiometer on the module, or set RAIN_ACTIVE_LOW to 0 in the sketch and upload again

## D. MATLAB live demo (15 min)
- [ ] MATLAB: cd to Review0_Prototype_Demo/matlab, type review0_dashboard
- [ ] Source: Serial (ESP), pick the COM port, baud 115200, Connect, then START live loop
- [ ] Move the sliders: sky (clear -> overcast), solar hour, wind; the pyranometer and PAR values and the angle change
- [ ] Wet the rain plate: mode becomes RAIN HARVEST, angle 3 degrees, cistern level rises
- [ ] Dry soil (probe in air): pump decision ON; probe in water: OFF
- [ ] Set wind to 18: WIND STOW 90 degrees
- [ ] Press Save evidence (PNG + CSV): files go to results/live_logs
- [ ] Screenshot/photo of the laptop with the real board connected (slide 14)

## E. Finish (before the review)
- [ ] Open ppt/Review0_Prototype_Demo.pptx, replace the three dashed frames on slide 14 with the photos/screenshots, fill the table with your numbers
- [ ] Supervisor signs the one-page write-up (documents/Review0_Supervisor_OnePager.docx, print it)
- [ ] Upload the PPT to the seminar-hall PC 10 minutes before the review

## If something fails
- No board / port busy: set Source = Simulate and say so. The dashboard, rules and results are identical.
- No MATLAB serial: use the Arduino Serial Monitor to show the live line, then Simulate for the decision logic.
"""
open(os.path.join(ROOT, 'documents', 'Review0_Morning_Checklist.md'), 'w', encoding='utf-8').write(chk)
os.makedirs(os.path.join(ROOT, 'test_data'), exist_ok=True)
with open(os.path.join(ROOT, 'test_data', 'bench_log_template.csv'), 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f); w.writerow(['test', 'reference', 'reference_value', 'measured_value', 'unit', 'criterion', 'pass(Y/N)', 'notes'])
    for row in [('DHT11 temperature', 'room thermometer', '', '', 'C', 'within 2 C', '', ''), ('DHT11 humidity', 'reference hygrometer', '', '', '% RH', 'within 5 % RH', '', ''), ('Soil raw, probe in air (dry)', '', '', '', 'ADC counts', 'high value', '', ''), ('Soil raw, probe in water (wet)', '', '', '', 'ADC counts', 'clearly lower than dry', '', ''),
                ('Rain detection time', 'drop on plate', '', '', 's', '<= 2 s', '', ''), ('Rain clear time', 'wipe plate', '', '', 's', '<= 30 s', '', ''), ('End to end: wet plate -> 3 degree mode in MATLAB', '', '', '', 's', '<= 5 s', '', '')]: w.writerow(row)

# ------------------------------------------------------------------ demo script and Q&A
qa = """# Review 0 - demo script and likely questions

## Marks to aim at (rubric)
Prototype demonstration 20 | Testing and validation evidence 10 | Feasibility and constraints 5 | SDG and cost mapping 5 | Time management 5 | Response to questions 5.

## 5-minute demo script
1. (30 s) "The goal is one decision: the panel angle that gives the most energy and still lets the crop receive its light."
2. (45 s) Slide 3 / dashboard: "These three are real sensors on the prototype: DHT11, soil moisture and rain. These three are simulated, shown as SIM: pyranometer, PAR and wind, because we do not have those sensors yet."
3. (60 s) Live: change sky from Clear to Overcast; GHI and PAR fall; the controller changes the angle. Move the solar hour; the angle follows.
4. (45 s) Wet the rain plate: "Rain sensor wet -> the panels go to a 3 degree slope, the cistern fills."
5. (30 s) Probe in air: soil % falls, pump ON from the harvested water. Probe in water: pump OFF.
6. (30 s) Wind slider to 18: stow at 90 degrees. "Safety overrides everything."
7. (60 s) Slides 9 to 12: "16 of 16 logic tests pass; on a clear day the controller keeps about 75 % of the best energy and 89 % of the best crop light, while pure sun tracking keeps 100 % energy and 27 % crop light."

## Likely questions and honest answers
1. **Is the pyranometer real?** No. It is computed from the sun model and a sky setting and is labelled SIM. The control code only reads a number, so a real pyranometer replaces that input without code changes.
2. **How do you know the virtual numbers are right?** The same equations were checked against PVGIS (annual PV within +5.0 %, of which +5.2 % is the sunshine dataset) and against Blender ray tracing (shading difference 0.010). They were not checked against a field instrument.
3. **What is "co-optimising"?** An objective J = w x energy + (1 - w) x crop light, each normalised to its best value at that moment; w is a slider. The optimiser tests every angle from 15 to 60 degrees in 0.5 degree steps.
4. **Why not just track the sun?** Tracking gives the most energy but leaves the crop about 27 % of its light target on a clear day, and the wire-drive design has two positions, not continuous tracking. It is shown only as a reference.
5. **Does the controller beat tracking?** Not on energy; it trades about a quarter of the energy for most of the crop light. It beats a fixed 45 degree angle on both. Which trade is right depends on w, which Review 4 prices.
6. **How accurate are the sensors?** DHT11: +-2 C and +-5 % RH (datasheet); our bench comparison is on slide 14. The resistive soil probe drifts and needs two-point recalibration each session.
7. **Why USB serial and not WiFi?** It needs no network and is reliable for a demo; the same protocol can go over WiFi/MQTT later.
8. **Why two boards?** The ESP8266 has one analog pin; the ESP32 is a backup with more ADC pins. Same serial lines, same MATLAB code.
9. **What does VPD do here?** It is computed from the DHT11 temperature and humidity. High VPD (> 2.5 kPa) or dry soil raises a stress index that lowers the light the crops need by up to 30 %, so the controller can accept more shade and more energy.
10. **What if the sensor fails?** A failed DHT11 read (-999) is ignored and the last good value is kept; a lost serial link is reported on the dashboard and Simulate mode takes over.
11. **Is the tilt real?** No actuator yet. The angle is shown on the MATLAB drawing; a motor driver with an encoder is the next step (Reviews 1 and 4 contain the full-scale wire-drive design).
12. **SDGs?** 7 (clean energy), 2 (food), 6 (rain harvest and efficient irrigation), 13 (climate action), 1/8 (monthly farmer income). Indicators are on slide 16.
13. **Cost?** Kit of a few hundred to about one thousand rupees at hobby prices (indicative, replace with the bill); a calibrated pyranometer is far more expensive, which is why it is simulated now.
14. **What has been validated and what not?** Validated against PVGIS and Blender; logic tested in software; sensors bench-tested today. Not validated against field data or a real pyranometer.
15. **What is next?** Real pyranometer and PAR sensor, a dead-band and motor driver, a week of logged data, comparison with local weather data.
"""
open(os.path.join(ROOT, 'documents', 'Review0_Demo_Script_and_QA.md'), 'w', encoding='utf-8').write(qa)
print('documents written')
